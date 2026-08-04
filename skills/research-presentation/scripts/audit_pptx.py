#!/usr/bin/env python3
"""Run lightweight, dependency-minimal structural checks on a research deck."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from pptx import Presentation


ANCHOR = re.compile(r"(?:fig(?:ure)?\.?\s*\d+|table\s*\d+|§\s*\d|p(?:age)?\.?\s*\d+|来源|source|arxiv|doi)", re.I)


def shape_text(shape) -> str:
    return getattr(shape, "text", "") or ""


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--rendered", type=Path)
    parser.add_argument("--report", type=Path, default=Path("qa-report.json"))
    args = parser.parse_args()
    prs = Presentation(args.pptx)
    width, height = prs.slide_width, prs.slide_height
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = [shape_text(shape) for shape in slide.shapes]
        nonempty = [text for text in texts if text.strip()]
        overflow = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                overflow.append(getattr(shape, "name", "unnamed"))
        slides.append({
            "slide": index,
            "has_title_like_text": bool(nonempty),
            "has_source_anchor": bool(ANCHOR.search(" ".join(nonempty))),
            "notes_present": slide.has_notes_slide,
            "shape_count": len(slide.shapes),
            "overflow_shapes": overflow,
        })

    rendered_count = None
    if args.rendered:
        rendered_count = len(list(args.rendered.glob("slide-*.png")))
    report = {
        "file": str(args.pptx.resolve()),
        "slide_count": len(prs.slides),
        "rendered_pngs": rendered_count,
        "all_slides_have_source_anchor": all(item["has_source_anchor"] for item in slides[1:-1]) if len(slides) > 2 else True,
        "all_slides_have_notes": all(item["notes_present"] for item in slides),
        "overflow_slide_count": sum(bool(item["overflow_shapes"]) for item in slides),
        "slides": slides,
    }
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in report if k != "slides"}, ensure_ascii=False, indent=2))
    return 0 if report["overflow_slide_count"] == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
